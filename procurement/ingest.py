"""Read the two supplied supplier ZIP formats without unpacking to disk."""

from __future__ import annotations

import io
import re
import zipfile
from dataclasses import dataclass
from datetime import date, datetime

import pandas as pd
from openpyxl import load_workbook


MONTHS = ('янв', 'фев', 'мар', 'апр', 'май', 'июн', 'июл', 'авг', 'сен', 'окт', 'ноя', 'дек')


def key(value):
    return str(value).strip() if value is not None else ''


def number(value):
    try:
        return float(str(value).replace(' ', '').replace(',', '.'))
    except (ValueError, TypeError):
        return None


def month_header(value):
    text = str(value or '').lower().strip()
    match = re.search(r'20\d{2}', text)
    if not match:
        return None
    for index, prefix in enumerate(MONTHS, 1):
        if prefix in text:
            return date(int(match.group()), index, 1)
    return None


@dataclass
class SupplierData:
    name: str
    monthly: pd.DataFrame
    events: pd.DataFrame
    catalog: pd.DataFrame
    incoming: pd.DataFrame
    as_of: date
    issues: list[str]


def _worksheet(archive, filename):
    payload = archive.read(filename)
    if len(payload) > 40_000_000:
        raise ValueError(f'Слишком большой файл Excel: {filename}')
    workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    try:
        yield from workbook.active.iter_rows(values_only=True)
    finally:
        workbook.close()


def _table(archive, filename):
    return list(_worksheet(archive, filename))


def _find(files, start):
    matches = [item for item in files if item.rsplit('/', 1)[-1].lower().startswith(start.lower())]
    if len(matches) != 1:
        raise ValueError(f'Ожидался ровно один файл «{start}*.xlsx», найдено: {len(matches)}')
    return matches[0]


def load_archive(content: bytes, supplier: str) -> SupplierData:
    if len(content) > 40_000_000:
        raise ValueError('ZIP-архив больше допустимых 40 МБ')
    is_iek = supplier == 'IEK'
    if supplier not in ('IEK', 'Systeme Electric'):
        raise ValueError('Неизвестный поставщик')
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        files = [x.filename for x in archive.infolist() if x.filename.lower().endswith('.xlsx')]
        if any(x.file_size > 40_000_000 for x in archive.infolist()):
            raise ValueError('В архиве есть слишком большой файл')
        sales = _table(archive, _find(files, 'Ежемесячные продажи'))
        stock = _table(archive, _find(files, 'Ежемесячные остатки'))
        moq = _table(archive, _find(files, 'MOQ'))
        txn = _table(archive, _find(files, 'Динамика'))
        transit = _table(archive, _find(files, 'Путь' if is_iek else 'Товар в пути'))

    issues = []
    products: dict[str, dict] = {}
    for row in moq[1:]:
        code = key(row[1 if is_iek else 2])
        if not code:
            continue
        pack = number(row[4])
        products[code] = dict(code=code, name=key(row[3 if is_iek else 1]),
                              article=key(row[2 if is_iek else 3]),
                              pack=int(pack) if pack and pack >= 1 and pack.is_integer() else None,
                              category='Не задана', balance=None, balance_source='', pending=0.0)

    s_header = stock[0]
    sept = [(i, v) for i, v in enumerate(s_header) if month_header(v) == date(2026, 9, 1)]
    if is_iek:
        if sept:
            for row in stock[3:]:
                code = key(row[2])
                if code in products:
                    val = number(row[sept[0][0]])
                    if val is not None:
                        products[code]['balance'] = max(0.0, val)
                        products[code]['balance_source'] = 'Остаток на начало сентября 2026; демонстрационный'
        issues.append('IEK: остаток за сентябрь указан на начало месяца; для рабочего заказа загрузите актуальный остаток.')
    else:
        issues.append('Systeme Electric: свободный остаток взят из сводной таблицы; дату и состав показателя нужно подтвердить.')

    monthly = []
    months = [(i, d) for i, cell in enumerate(sales[0]) if (d := month_header(cell))]
    for row in sales[2:]:
        code = key(row[1])
        if not code:
            continue
        if code not in products:
            products[code] = dict(code=code, name=key(row[0]), article='', pack=None,
                                  category='Не задана', balance=None, balance_source='', pending=0.0)
        for col, period in months:
            qty = number(row[col]) if col < len(row) else None
            monthly.append((code, period, max(0.0, qty or 0.0)))
    issues.append('Отрицательные помесячные продажи для прогноза ограничены нулём; первичные возвраты требуют сверки.')

    incoming = []
    if is_iek:
        for row in transit[1:]:
            code = key(row[0])
            if not code:
                continue
            for col, header in enumerate(transit[0][3:9], 3):
                qty = number(row[col])
                date_match = re.search(r'до\s*(\d{2}\.\d{2}\.\d{4})', str(header))
                if qty and qty > 0 and date_match:
                    incoming.append((code, datetime.strptime(date_match.group(1), '%d.%m.%Y').date(), qty))
    else:
        for row in transit[2:]:
            code = key(row[2])
            if not code:
                continue
            if code in products:
                products[code]['category'] = key(row[4]) or 'Не задана'
                free = number(row[51])
                if free is not None:
                    products[code]['balance'] = max(0.0, free)
                    products[code]['balance_source'] = 'Свободный остаток в сводной таблице'
                pending = number(row[54])
                products[code]['pending'] = max(0.0, pending or 0.0)
        issues.append('Systeme Electric: для товара в пути нет подтверждённой даты прибытия; он не вычитается из заказа.')

    events = []
    dates = []
    for row in txn[1:]:
        code = key(row[3])
        raw_date = row[0]
        qty = number(row[7]) if len(row) > 7 else None
        if not code or qty is None:
            continue
        try:
            when = datetime.strptime(str(raw_date)[:19], '%d.%m.%Y %H:%M:%S').date()
        except ValueError:
            continue
        if not key(row[2]).startswith('Расходная накладная'):
            continue
        dates.append(when)
        if qty > 0:
            events.append((code, when, key(row[1]), qty))
    if not dates:
        raise ValueError('В архиве нет распознанной истории продаж')
    issues.append('В выгрузке нет ID клиента; крупные разовые продажи выявляются по документу и артикулу.')
    issues.append('Журнала stockout в архиве нет; компенсация упущенного спроса включается после загрузки его CSV.')
    return SupplierData(supplier, pd.DataFrame(monthly, columns=['code', 'month', 'qty']),
                        pd.DataFrame(events, columns=['code', 'date', 'document', 'qty']),
                        pd.DataFrame(products.values()),
                        pd.DataFrame(incoming, columns=['code', 'eta', 'qty']), max(dates), issues)
